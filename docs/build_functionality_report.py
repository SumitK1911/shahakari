from __future__ import annotations

from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "Sahakari_Functionality_Report.docx"


BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
INK = "172033"
MUTED = "617184"
LIGHT_BLUE = "E8EEF5"
LIGHT_GRAY = "F2F4F7"
PALE_GREEN = "EEF8F1"
GREEN = "126C43"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_text(cell, text: str, bold: bool = False, color: str = INK) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.size = Pt(9.5)
    run.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_table_geometry(table, widths: list[float]) -> None:
    table.autofit = False
    for row in table.rows:
        for idx, width in enumerate(widths):
            row.cells[idx].width = Inches(width)
            tc_pr = row.cells[idx]._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(int(width * 1440)))
            tc_w.set(qn("w:type"), "dxa")


def add_table(doc: Document, headers: list[str], rows: list[list[str]], widths: list[float]) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    set_table_geometry(table, widths)
    hdr = table.rows[0]
    for idx, header in enumerate(headers):
        set_cell_text(hdr.cells[idx], header, bold=True, color=DARK_BLUE)
        set_cell_shading(hdr.cells[idx], LIGHT_BLUE)
    for row in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            set_cell_text(cells[idx], value)
    doc.add_paragraph()


def add_heading(doc: Document, text: str, level: int = 1) -> None:
    paragraph = doc.add_heading(text, level=level)
    for run in paragraph.runs:
        run.font.color.rgb = RGBColor.from_string(BLUE if level < 3 else DARK_BLUE)


def add_bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        paragraph = doc.add_paragraph(style="List Bullet")
        paragraph.paragraph_format.space_after = Pt(4)
        run = paragraph.add_run(item)
        run.font.size = Pt(10.5)
        run.font.color.rgb = RGBColor.from_string(INK)


def add_callout(doc: Document, title: str, body: str, fill: str = PALE_GREEN) -> None:
    table = doc.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    set_table_geometry(table, [6.5])
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(title)
    r.bold = True
    r.font.size = Pt(11)
    r.font.color.rgb = RGBColor.from_string(GREEN)
    p2 = cell.add_paragraph()
    p2.paragraph_format.space_after = Pt(0)
    r2 = p2.add_run(body)
    r2.font.size = Pt(10)
    r2.font.color.rgb = RGBColor.from_string(INK)
    doc.add_paragraph()


def style_document(doc: Document) -> None:
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    for name, size, color, before, after in [
        ("Heading 1", 16, BLUE, 18, 10),
        ("Heading 2", 13, BLUE, 14, 7),
        ("Heading 3", 12, DARK_BLUE, 10, 5),
    ]:
        style = styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.font.bold = True
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)


def add_cover(doc: Document) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run("Sahakari Finance Application")
    r.font.size = Pt(26)
    r.font.color.rgb = RGBColor.from_string(INK)
    r.bold = True

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(18)
    r = p.add_run("Functionality Report and Comparison Checklist")
    r.font.size = Pt(14)
    r.font.color.rgb = RGBColor.from_string(MUTED)

    add_callout(
        doc,
        "Purpose",
        "This document summarizes the functionality currently implemented in the Sahakari cooperative finance application, how key workflows operate, what automation and reporting exist, and which production-grade features still need comparison against a real cooperative organization.",
    )

    add_table(
        doc,
        ["Field", "Value"],
        [
            ["Report date", date.today().isoformat()],
            ["Application type", "Local-network cooperative finance / sahakari operations system"],
            ["Frontend", "React + Vite desktop operations UI"],
            ["Backend", "FastAPI, SQLAlchemy, PostgreSQL-ready models, JWT/RBAC dependencies"],
            ["Current status", "Functional foundation with expanded operations UI; not yet complete for live regulated production"],
        ],
        [1.7, 4.8],
    )
    doc.add_page_break()


def add_section_summary(doc: Document) -> None:
    add_heading(doc, "1. Executive Summary")
    add_callout(
        doc,
        "Current position",
        "The app is no longer only a small registration form. It now includes member onboarding, scalable member search, member 360 profile, shares, savings, cash counter, loans, repayments, accounting, rules, automation, reports, approval queues, Excel-ready export, security/compliance screens, and dashboard metrics.",
    )
    add_bullets(
        doc,
        [
            "The system is suitable as a working prototype/foundation for a cooperative finance office.",
            "The backend already includes financial entities such as members, nominees, shares, savings accounts, loans, installments, payments, accounts, journal entries, ledger entries, fiscal closings, scheduler jobs, calculation results, and audit logs.",
            "The frontend exposes operational workspaces instead of only a registration form.",
            "The biggest remaining production gaps are full maker-checker approval, locked-period enforcement, richer product masters, true XLSX import/export, backup/restore UI, branch/cashier session depth, and formal financial statement generation.",
        ],
    )


def add_module_matrix(doc: Document) -> None:
    add_heading(doc, "2. Implemented Functionality Matrix")
    rows = [
        ["Authentication", "Staff login using email/password and JWT token storage.", "Implemented"],
        ["Dashboard", "Members, savings, loan exposure, risk, cash movement, pending KYC, high-risk cases.", "Implemented"],
        ["Member onboarding", "Detailed person form: KYC, address, documents, family, guardian, nominee, occupation, income, risk flags, media references.", "Implemented"],
        ["Member search", "Server-side search by member number, name, or phone with skip/limit paging and total count.", "Implemented"],
        ["Member 360 profile", "Member data, nominees, shares, share activity, savings accounts, savings ledger, loans, schedules, payments, calculations.", "Implemented"],
        ["Share capital", "Share purchase posting and share register display.", "Implemented"],
        ["Savings accounts", "Open savings account, list active balances, post deposit and withdrawal.", "Implemented"],
        ["Cash counter", "Deposit and withdrawal screens connected to savings transactions and ledger posting.", "Implemented"],
        ["Loans", "Loan disbursement, flat installment schedule generation, loan portfolio, schedule view, installment payment.", "Implemented"],
        ["Accounting", "Chart of accounts and trial balance endpoints/UI.", "Implemented"],
        ["Rules engine", "Interest, penalty, accrual rule CRUD/read views; calculation history.", "Implemented foundation"],
        ["Automation", "Manual automation run, daily interest trigger, overdue marking trigger, scheduler job list, calculation results.", "Implemented foundation"],
        ["Reports", "Portfolio summary, daily operations, approval queue, high-risk queue, overdue queue, trial balance.", "Implemented foundation"],
        ["Excel/export", "CSV member portfolio export, openable in Excel.", "Implemented basic"],
        ["Security/compliance", "RBAC dependencies, audit logging in services, security workspace, risk flags.", "Implemented foundation"],
        ["Closings", "Day/month/year closing models and closing register display.", "Foundation only"],
    ]
    add_table(doc, ["Module", "What it does now", "Status"], rows, [1.45, 4.35, 1.0])


def add_workflows(doc: Document) -> None:
    add_heading(doc, "3. How Key Workflows Work")
    workflows = [
        ["Member onboarding", "Staff enters personal, KYC, document, address, family, nominee, occupation, income and risk information. The backend stores the member, JSON profile sections, nominee records and audit entry."],
        ["Finding one member among 10k+", "Staff searches by member number, name or phone. The backend applies indexed search and returns paged results. Staff clicks Open to load the member 360 profile instead of scrolling thousands of rows."],
        ["Member 360 review", "The profile endpoint gathers the selected member's shares, savings accounts, savings transactions, loans, installments, payments, nominees and calculation results into one operational view."],
        ["Savings deposit", "Staff selects account, enters amount/narration, posts deposit. Backend locks the account, increases balance, creates savings transaction, posts double-entry ledger lines, and writes audit log."],
        ["Savings withdrawal", "Staff selects account, enters amount/narration, posts withdrawal. Backend checks active account and sufficient balance, decreases balance, posts transaction, ledger and audit."],
        ["Loan disbursement", "Staff selects member and loan facts. Backend creates active loan, generates flat monthly installments, posts loan disbursement ledger entry and audit log."],
        ["Loan repayment", "Staff opens loan schedule and pays an installment. Backend marks installment paid, reduces outstanding principal, posts principal and interest ledger entries, and audits the payment."],
        ["Automation run", "Admin/manager runs automation for a date. Backend calculates savings interest and overdue penalties through the rules engine, updates live records, stores calculation results and returns posting summary."],
        ["Reports/export", "Reports screen loads portfolio totals, queues and trial balance. CSV export downloads member portfolio data for Excel review."],
    ]
    add_table(doc, ["Workflow", "Operational behavior"], workflows, [1.8, 4.7])


def add_api_frontend(doc: Document) -> None:
    add_heading(doc, "4. Current Screens and API Coverage")
    add_heading(doc, "Frontend Workspaces", level=2)
    add_table(
        doc,
        ["Screen", "Main capability"],
        [
            ["Dashboard", "Daily snapshot, quick actions, recent members, health signals."],
            ["Members", "Onboarding, search/paging, member 360 profile."],
            ["Shares", "Share purchase and share register."],
            ["Savings", "Open account and view account balances."],
            ["Cash Counter", "Post deposit/withdrawal and monitor counter balances."],
            ["Loans", "Disburse loan, view portfolio, open schedule, post repayment."],
            ["Reports", "Portfolio, approval queue, high-risk cases, overdue queue, trial balance, CSV export."],
            ["Automation", "Daily/monthly/fiscal automation buttons, scheduler jobs, closing register, calculations."],
            ["Rules", "Create interest rule and inspect interest/penalty/accrual rules."],
            ["Accounting", "Chart of accounts and trial balance."],
            ["Security", "Compliance controls, RBAC/audit/export/backup guidance."],
        ],
        [1.6, 4.9],
    )
    add_heading(doc, "Backend API Groups", level=2)
    add_bullets(
        doc,
        [
            "/auth: login",
            "/dashboard: operational metrics",
            "/branches: branch list/create",
            "/members: list, count, profile, create, update, nominees",
            "/shares: list and purchase",
            "/savings: accounts, deposit, withdrawal",
            "/loans: list, create, installments, payments",
            "/accounting: chart of accounts, trial balance",
            "/reports: portfolio summary, daily operations, approval queue, CSV export",
            "/automation: manual run and async triggers",
            "/rules: interest, penalty, accrual, scheduler jobs, calculation results",
        ],
    )


def add_calculations(doc: Document) -> None:
    add_heading(doc, "5. Calculations, Ledgers and Automation")
    add_table(
        doc,
        ["Area", "Current behavior", "Notes"],
        [
            ["Savings interest", "Rules engine checks active interest rule by savings product/account type; falls back to account interest rate.", "Foundation implemented; product-level configuration should be expanded."],
            ["Penalty", "Rules engine calculates overdue installment penalty from configured rule or fallback rate.", "Supports fixed/percentage/compound concepts in rule schema."],
            ["Loan schedule", "Flat monthly principal plus monthly interest schedule generated at disbursement.", "EMI/reducing/diminishing variants are not yet fully implemented."],
            ["Double-entry ledger", "Savings, shares, loan disbursement and repayment services post journal/ledger entries.", "Needs stricter balancing audits and approval before live use."],
            ["Calculation records", "Automation stores calculation_results for traceability.", "Good foundation for explainable historical calculations."],
            ["Closings", "Models exist for day, month and year closings; UI displays closing register.", "Actual close/lock workflow needs more implementation."],
        ],
        [1.45, 3.25, 1.8],
    )


def add_excel(doc: Document) -> None:
    add_heading(doc, "6. Excel and Reporting")
    add_callout(
        doc,
        "Excel recommendation",
        "Excel should be used for exports, imports, migration, audit sharing and offline review. It should not be used as the primary database for live cooperative finance records.",
        fill=LIGHT_GRAY,
    )
    add_bullets(
        doc,
        [
            "Available now: CSV member portfolio export at /api/v1/reports/portfolio-export.csv.",
            "CSV can be opened in Microsoft Excel and used for comparison or audit review.",
            "Recommended next step: true .xlsx exports with multiple sheets for members, savings, loans, shares, trial balance and daily cash.",
            "Recommended next step: bulk import with validation preview for opening balances and member migration.",
        ],
    )


def add_missing_checklist(doc: Document) -> None:
    add_heading(doc, "7. Production Gap Checklist")
    rows = [
        ["Maker-checker approval", "Formal approval screens for KYC, loan, withdrawal, adjustment, reversal, closing.", "Missing"],
        ["Locked-period enforcement", "Reject edits/postings after day/month/year close unless reversal workflow is used.", "Mostly missing"],
        ["Product masters", "Savings products, FD products, loan products, share products, fee products.", "Missing"],
        ["Advanced loan engine", "EMI, reducing balance, diminishing, partial payment, reschedule, collateral, guarantors.", "Partial"],
        ["Cashier session", "Opening cash, denomination count, handover, vault transfer, shortage/excess handling.", "Partial"],
        ["Backup/restore UI", "Encrypted backups, restore validation, backup schedule and status dashboard.", "Missing"],
        ["True XLSX reports", "Formatted multi-sheet Excel export and import validation.", "Basic CSV only"],
        ["Financial statements", "Balance sheet, profit/loss, ledger statement, cash book, member statement, regulatory reports.", "Partial"],
        ["Notifications", "SMS/email/reminder outbox for due loans, overdue, maturity, KYC expiry.", "Missing"],
        ["Session/rate limit", "Refresh tokens, session tracking, login rate limiting.", "Partial"],
        ["Branch scope", "Strict branch-scoped data access for all modules.", "Partial"],
        ["Audit viewer", "Searchable audit log UI with filters and export.", "Missing UI"],
        ["Backup of documents/media", "Real file upload/storage for photo, signature, thumbprint and KYC files.", "Reference fields only"],
    ]
    add_table(doc, ["Feature", "Expected in real production", "Current"], rows, [1.65, 4.1, 0.75])


def add_comparison_template(doc: Document) -> None:
    add_heading(doc, "8. Real-Organization Comparison Template")
    add_table(
        doc,
        ["Area to compare", "Questions for the cooperative office", "Mark"],
        [
            ["Membership", "What exact KYC fields, documents, nominees, guardian rules and membership types are required?", ""],
            ["Savings", "Which account products exist, interest methods, withdrawal limits, dormancy rules and fees?", ""],
            ["Loans", "What loan products, approval stages, guarantors, collateral, interest methods and penalty rules are used?", ""],
            ["Shares", "Share purchase limits, transfer/refund rules, certificate format and dividend process?", ""],
            ["Cash counter", "How many counters, cashier opening/closing process, vault movement and denomination requirements?", ""],
            ["Accounting", "Chart of accounts, voucher approval, fiscal year close, reserves and dividend postings?", ""],
            ["Reports", "Daily, monthly, annual, regulatory, board and member statement reports required?", ""],
            ["Excel", "Which data must be exported/imported, and what format does the office currently use?", ""],
            ["Security", "Roles, permissions, branch access, audit review and password/session policies?", ""],
            ["Backup", "Backup location, frequency, encryption, external drive/NAS and restore testing policy?", ""],
        ],
        [1.45, 4.55, 0.5],
    )


def add_recommended_next_steps(doc: Document) -> None:
    add_heading(doc, "9. Recommended Next Development Phases")
    add_table(
        doc,
        ["Phase", "Goal", "Deliverables"],
        [
            ["Phase 1", "Approval and control", "Maker-checker queues, approval actions, locked-period checks, reversal workflow."],
            ["Phase 2", "Product depth", "Savings/FD/loan product masters, interest methods, fees, collateral and guarantors."],
            ["Phase 3", "Cash operations", "Cashier sessions, vault movement, denominations, handover and day close."],
            ["Phase 4", "Reports and Excel", "Member statement, passbook, cash book, balance sheet, P&L, XLSX export/import."],
            ["Phase 5", "Backup and production hardening", "Encrypted backups, restore UI, audit viewer, session/rate-limit, branch scope tests."],
        ],
        [1.0, 2.0, 3.5],
    )


def add_footer(doc: Document) -> None:
    section = doc.sections[0]
    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = p.add_run("Sahakari Functionality Report")
    r.font.size = Pt(8)
    r.font.color.rgb = RGBColor.from_string(MUTED)


def build() -> None:
    doc = Document()
    style_document(doc)
    add_footer(doc)
    add_cover(doc)
    add_section_summary(doc)
    add_module_matrix(doc)
    add_workflows(doc)
    add_api_frontend(doc)
    add_calculations(doc)
    add_excel(doc)
    add_missing_checklist(doc)
    add_comparison_template(doc)
    add_recommended_next_steps(doc)
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
